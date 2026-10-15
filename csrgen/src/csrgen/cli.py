import argparse
from pathlib import Path
import sys
import warnings
from zipfile import BadZipFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from .diagnostics import NumericCellWarning, ReservedFieldWarning, TableError
from .layout import layout_from_overrides
from .naming import resolve_output_path
from .pipeline import GenerationOptions, run_pipeline
from .reader import ReadOptions


def write_template(path):
    """Produce the screenshot's B:G layout, including its blank separator row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sheet1"
    rows = [
        [None, "Offset(hex)", "signal", "Rst Value(hexadecimal)", "bit", "access", "description"],
        [None, "0x0", "a", "1", "0", "rw", "Control enable"],
        [None, None, "b", "1", "2:1", "ro", "External status"],
        [None, None, "c", "0", "30:3", "rw1c", "Sticky status, write one to clear"],
        [],
        [None, "0x4", "a1", "1", "0", "rw", "Control enable"],
        [None, None, "b1", "1", "2:1", "ro", "External status"],
        [None, None, "c1", "0", "30:3", "rw1c", "Sticky status, write one to clear"],
        [None, None, "reserved", "0", "31", "ro", "Reads zero"],
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1][1:]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for column, width in {"A": 3, "B": 16, "C": 20, "D": 28, "E": 12, "F": 12, "G": 44}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B2"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(path)
    finally:
        workbook.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description="Generate APB3 register SystemVerilog from XLSX (32-bit data, no strb).")
    parser.add_argument("xlsx", nargs="?", type=Path, help="input register table")
    parser.add_argument("--sheet", "-s", default="sheet1", help="worksheet name (default: sheet1, case insensitive fallback)")
    parser.add_argument("--aw", "-aw", type=int, default=16, help="address width, 2..32 (default: 16)")
    parser.add_argument("--header-rows", type=int, default=1,
                        help="rows to skip before register data; 0 for no header (default: 1)")
    parser.add_argument("--column", action="append", default=[], metavar="FIELD=COLUMN",
                        help="override a template column, e.g. offset=A; letters or 1-based numbers")
    parser.add_argument("--lint", action="store_true", help="review RTL with Verilator before writing output")
    parser.add_argument("--lint-tool", default="verilator", help="Verilator executable name or path")
    parser.add_argument("--format", action="store_true", help="format RTL with Verible before writing output")
    parser.add_argument("--format-tool", default="verible-verilog-format", help="Verible formatter executable name or path")
    parser.add_argument("--module", "-m", help="base RTL module name (default: input filename stem)")
    parser.add_argument("--module-prefix", help="prepend <prefix>_ to the base module name; does not change the filename")
    output_group = parser.add_mutually_exclusive_group()
    output_group.add_argument("--output", "-o", type=Path, help="output filename and path (default: ./<input-stem>.sv)")
    output_group.add_argument("--file-prefix", help="filename: <prefix>_<input-stem>.sv; does not change the module name")
    parser.add_argument("--template", type=Path, metavar="XLSX", help="create an example XLSX table and exit")
    args = parser.parse_args(argv)
    if args.template and args.xlsx:
        parser.error("--template cannot be combined with an input XLSX")
    if not args.template and args.xlsx is None:
        parser.error("input XLSX is required (or use --template example.xlsx)")
    try:
        if args.template:
            write_template(args.template)
            print(f"Created template: {args.template.resolve()}")
            return 0
        output = resolve_output_path(args.xlsx, args.output, args.file_prefix)
        options = GenerationOptions(
            read=ReadOptions(sheet=args.sheet),
            layout=layout_from_overrides(args.header_rows, args.column),
            module=args.module, module_prefix=args.module_prefix, aw=args.aw,
            lint=args.lint, format=args.format,
            lint_tool=args.lint_tool, format_tool=args.format_tool,
        )
        with warnings.catch_warnings(record=True) as notices:
            warnings.simplefilter("always", ReservedFieldWarning)
            warnings.simplefilter("always", NumericCellWarning)
            result = run_pipeline(args.xlsx, options)
        for notice in notices:
            print(f"csrgen: warning: {notice.message}", file=sys.stderr)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(result.rtl, encoding="utf-8")
        if result.reviews:
            print(f"RTL review passed ({len(result.reviews)} checks)")
        if args.format:
            print("RTL formatting applied")
        count = sum(not item.reserved for reg in result.registers for item in reg.fields)
        print(f"Generated {output.resolve()} ({len(result.registers)} registers, {count} fields, aw={args.aw})")
        return 0
    except (TableError, OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        print(f"csrgen: error: {exc}", file=sys.stderr)
        return 2
