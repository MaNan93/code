"""Parse field tables and generate synthesizable SystemVerilog APB3 slaves."""
from dataclasses import dataclass, field, replace
import math
from pathlib import Path
import re
import warnings

from openpyxl import load_workbook


class TableError(ValueError):
    """Invalid register table or generator configuration."""


class ReservedFieldWarning(UserWarning):
    """Reserved field metadata is ignored in favor of constant zero / RO."""


class NumericCellWarning(UserWarning):
    """A numeric Excel cell was used as decimal where the column is hexadecimal."""


# Verilog/SystemVerilog reserved words cannot be used as module/port identifiers.
KEYWORDS = set("""
accept_on alias always always_comb always_ff always_latch and assert assign assume
automatic before begin bind bins binsof bit break buf bufif0 bufif1 byte case casex
casez cell chandle checker class clocking cmos config const constraint context
continue cover covergroup coverpoint cross deassign default defparam design disable
dist do edge else end endcase endchecker endclass endclocking endconfig endfunction
endgenerate endgroup endinterface endmodule endpackage endprimitive endprogram
endproperty endsequence endspecify endtable endtask enum event eventually expect
export extends extern final first_match for force foreach forever fork forkjoin
function generate genvar global highz0 highz1 if iff ifnone ignore_bins illegal_bins
implements implies import incdir include initial inout input inside instance int
integer interconnect interface intersect join join_any join_none large let liblist
library local localparam logic longint macromodule matches medium modport module
nand negedge nettype new nexttime nmos nor noshowcancelled not notif0 notif1 null or
output package packed parameter pmos posedge primitive priority program property
protected pull0 pull1 pulldown pullup pulsestyle_ondetect pulsestyle_onevent pure
rand randc randcase randsequence rcmos real realtime ref reg reject_on release
repeat restrict return rnmos rpmos rtran rtranif0 rtranif1 s_always s_eventually
s_nexttime s_until s_until_with scalared sequence shortint shortreal showcancelled
signed small soft solve specify specparam static string strong strong0 strong1
struct super supply0 supply1 sync_accept_on sync_reject_on table tagged task this
throughout time timeprecision timeunit tran tranif0 tranif1 tri tri0 tri1 triand
trior trireg type typedef union unique unique0 unsigned until until_with untyped use
uwire var vectored virtual void wait wait_order wand weak weak0 weak1 while wildcard
wire with within wor xnor xor
""".split())
INTERNAL_NAMES = {"pclk", "presetn", "psel", "penable", "pwrite", "paddr",
                  "pwdata", "prdata", "pready", "pslverr", "AW",
                  "csr_addr_hit", "csr_write", "csr_read", "csr_data"}


@dataclass(frozen=True)
class Field:
    name: str
    msb: int
    lsb: int
    reset: int
    access: str
    description: str = ""
    row: int = 0

    @property
    def width(self):
        return self.msb - self.lsb + 1

    @property
    def reserved(self):
        return self.name.lower() == "reserved"

    @property
    def bits(self):
        return f"[{self.msb}:{self.lsb}]" if self.width > 1 else f"[{self.lsb}]"


@dataclass
class Register:
    offset: int
    fields: list[Field] = field(default_factory=list)


def validate_identifier(name, label="signal"):
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_$]*", name) or name in KEYWORDS:
        raise TableError(f"{label}: invalid Verilog identifier {name!r}")


def validate_aw(aw):
    if isinstance(aw, bool) or not isinstance(aw, int) or not 2 <= aw <= 32:
        raise TableError("aw must be an integer between 2 and 32")


def _empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


TEMPLATE_COLUMNS = ("offset", "signal", "reset", "bit", "access", "description")


def _integer(value, base, label):
    if isinstance(value, bool):
        raise TableError(f"{label}: boolean values are not numbers")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (not math.isfinite(value) or not value.is_integer()):
            raise TableError(f"{label}: expected an integer, got {value!r}")
        # Excel numeric cells represent numbers, not hexadecimal digit strings.
        return int(value)
    token = str(value).strip().replace("_", "")
    if token.startswith("="):
        raise TableError(f"{label}: formulas are not supported; use a literal value")
    try:
        return int(token, base)
    except ValueError as exc:
        raise TableError(f"{label}: invalid number {value!r}") from exc


def _hex_cell(value, label):
    """Parse an offset/reset cell; warn when a numeric cell reads differently as hex."""
    number = _integer(value, 16, label)
    if not isinstance(value, str) and abs(number) >= 10:
        warnings.warn(
            f"{label}: numeric cell {number} is used as decimal (0x{number:X}); "
            f"enter it as text for hexadecimal 0x{number}",
            NumericCellWarning, stacklevel=3,
        )
    return number


def _normalize_reserved(item, offset):
    if not item.reserved:
        return item
    label = f"row {item.row}, offset 0x{offset:X}, signal {item.name!r}"
    try:
        reset_ok = _integer(item.reset, 16, label + " reset") == 0
    except TableError:
        reset_ok = False
    if not reset_ok or item.access != "ro":
        warnings.warn(
            f"{label}: reserved metadata reset={item.reset!r}, access={item.access!r}; "
            "using reset=0, access=ro (tie0, writes ignored)",
            ReservedFieldWarning, stacklevel=2,
        )
    return replace(item, reset=0, access="ro")


def _reserved_ranges(reg):
    """Return explicit reserved fields and all gaps as contiguous bit ranges."""
    defined = 0
    for item in reg.fields:
        if not item.reserved:
            defined |= ((1 << item.width) - 1) << item.lsb
    ranges = []
    bit = 31
    while bit >= 0:
        if defined & (1 << bit):
            bit -= 1
            continue
        msb = bit
        while bit >= 0 and not defined & (1 << bit):
            bit -= 1
        ranges.append((msb, bit + 1))
    return ranges


def validate_registers(registers, aw):
    validate_aw(aw)
    if not registers:
        raise TableError("no register fields found")
    offsets = set()
    names = set(INTERNAL_NAMES)
    names.update(f"OFFSET_0x{reg.offset:x}" for reg in registers)
    names.update(f"csr_reg_data_0x{reg.offset:x}" for reg in registers)
    for reg in registers:
        if reg.offset < 0 or reg.offset % 4:
            raise TableError(f"offset 0x{reg.offset:X}: expected a nonnegative, 4-byte aligned address")
        if reg.offset + 3 >= 1 << aw:
            raise TableError(f"offset 0x{reg.offset:X}: exceeds aw={aw}")
        if reg.offset in offsets:
            raise TableError(f"duplicate register offset 0x{reg.offset:X}")
        offsets.add(reg.offset)
        if not reg.fields:
            raise TableError(f"offset 0x{reg.offset:X}: no fields")
        used = 0
        for item in reg.fields:
            item = _normalize_reserved(item, reg.offset)
            label = f"row {item.row}, offset 0x{reg.offset:X}, signal {item.name!r}"
            if not 0 <= item.lsb <= item.msb <= 31:
                raise TableError(f"{label}: bits must satisfy 0 <= lsb <= msb <= 31")
            if item.access not in {"rw", "ro", "rw1c", "rwsc"}:
                raise TableError(f"{label}: unsupported access {item.access!r}; use rw, ro, rw1c, rwsc")
            if not 0 <= item.reset < 1 << item.width:
                raise TableError(f"{label}: reset value 0x{item.reset:X} does not fit {item.width} bits")
            mask = ((1 << item.width) - 1) << item.lsb
            if used & mask:
                raise TableError(f"{label}: overlapping bit fields")
            used |= mask
            if item.reserved:
                continue
            validate_identifier(item.name, label)
            ports = [item.name]
            if item.access == "rw1c":
                ports.append(item.name + "_hwset")
            for name in ports:
                if name in names:
                    raise TableError(f"{label}: duplicate or conflicting port name {name!r}")
                names.add(name)


def read_xlsx(path, sheet="sheet1", aw=16):
    """Read template columns B:G from row 2; ignore row 1's header text."""
    validate_aw(aw)
    workbook = load_workbook(Path(path), read_only=True, data_only=False)
    try:
        exact = [name for name in workbook.sheetnames if name == sheet]
        matches = exact or [name for name in workbook.sheetnames if name.casefold() == sheet.casefold()]
        if len(matches) != 1:
            raise TableError(f"sheet {sheet!r} not found or ambiguous; available: {', '.join(workbook.sheetnames)}")
        rows = enumerate(
            workbook[matches[0]].iter_rows(min_row=2, min_col=2, max_col=7, values_only=True),
            2,
        )
        registers = {}
        current_offset = None
        offset_rows = {}
        for row_num, values in rows:
            cells = dict(zip(TEMPLATE_COLUMNS, values))
            if all(_empty(value) for value in cells.values()):
                continue
            label = f"row {row_num}"
            if not _empty(cells["offset"]):
                offset = _hex_cell(cells["offset"], label + " offset")
                # Repeating the offset on consecutive rows is fine; returning to an
                # earlier one would silently merge two register blocks.
                if offset != current_offset and offset in offset_rows:
                    raise TableError(f"{label}: duplicate register offset 0x{offset:X}, "
                                     f"already defined at row {offset_rows[offset]}")
                offset_rows.setdefault(offset, row_num)
                current_offset = offset
            if _empty(cells["signal"]):
                if any(not _empty(cells[key]) for key in ("reset", "bit", "access")):
                    raise TableError(f"{label}: signal is missing")
                # Permit a row containing only a register offset / description.
                continue
            if current_offset is None:
                raise TableError(f"{label}: first field requires an offset")
            name = str(cells["signal"]).strip()
            reserved = name.lower() == "reserved"
            required = ("bit",) if reserved else ("reset", "bit", "access")
            for key in required:
                if _empty(cells[key]):
                    raise TableError(f"{label}: {key} is missing")
            bit = str(cells["bit"]).strip()
            # Excel reads an integer-valued numeric cell as int (occasionally float).
            if isinstance(cells["bit"], (int, float)) and not isinstance(cells["bit"], bool):
                bit = str(_integer(cells["bit"], 10, label + " bit"))
            match = re.fullmatch(r"\[?\s*(\d+)\s*(?::\s*(\d+)\s*)?\]?", bit)
            if not match:
                raise TableError(f"{label}: invalid bit range {cells['bit']!r}; use 0 or 30:3")
            msb = int(match[1])
            lsb = int(match[2]) if match[2] is not None else msb
            item = Field(
                name, msb, lsb,
                cells["reset"] if reserved else _hex_cell(cells["reset"], label + " reset"),
                str(cells["access"] or "").strip().lower(),
                str(cells.get("description") or "").strip(), row_num,
            )
            item = _normalize_reserved(item, current_offset)
            reg = registers.setdefault(current_offset, Register(current_offset))
            reg.fields.append(item)
        result = sorted(registers.values(), key=lambda reg: reg.offset)
        validate_registers(result, aw)
        return result
    finally:
        workbook.close()


def _literal(item):
    return f"{item.width}'h{item.reset:X}"


def generate_rtl(registers, module="apb_csr", aw=16):
    validate_identifier(module, "module name")
    validate_registers(registers, aw)
    fields = [item for reg in registers for item in reg.fields if not item.reserved]
    ports = [("input", "", "pclk"), ("input", "", "presetn"),
             ("input", "", "psel"), ("input", "", "penable"),
             ("input", "", "pwrite"), ("input", "[AW-1:0]", "paddr"),
             ("input", "[31:0]", "pwdata"), ("output", "[31:0]", "prdata"),
             ("output", "", "pready"), ("output", "", "pslverr")]
    for item in fields:
        direction = "input" if item.access == "ro" else "output"
        width = f"[{item.width-1}:0]" if item.width > 1 else ""
        ports.append((direction, width, item.name))
        if item.access == "rw1c":
            ports.append(("input", width, item.name + "_hwset"))
    lines = ["// Generated by csrgen. APB3, 32-bit data, no byte strobes.",
             "// rw1c: hardware set wins over software clear.",
             f"module {module} #(", f"    parameter integer AW = {aw}", ") (",
             ",\n".join(f"    {direction:<7}logic {width:<10}{name}"
                        for direction, width, name in ports), ");", ""]
    for reg in registers:
        lines.append(f"    localparam logic [AW-1:0] OFFSET_0x{reg.offset:x} = AW'(32'h{reg.offset:08X});")
    lines += ["", "    logic csr_addr_hit;", "    logic [31:0] csr_data;",
              "    logic csr_write;", "    logic csr_read;"]
    lines.extend(f"    logic [31:0] csr_reg_data_0x{reg.offset:x}; // offset 0x{reg.offset:X}"
                 for reg in registers)
    lines += ["", "    assign pready = 1'b1;",
              "    assign pslverr = psel && penable && !csr_addr_hit;",
              "    assign csr_write = psel && penable && pwrite && pready && !pslverr;",
              "    // Capture read data at the end of SETUP, before ACCESS completes.",
              "    assign csr_read = psel && !penable && !pwrite && csr_addr_hit;",
              "", "    always_comb begin", "        csr_addr_hit = 1'b0;",
              "        case (paddr)"]
    for reg in registers:
        lines.append(f"            OFFSET_0x{reg.offset:x}: csr_addr_hit = 1'b1;")
    lines += ["            default: csr_addr_hit = 1'b0;", "        endcase", "    end",
              "", "    always_comb begin", "        csr_data = 32'b0;",
              "        if (csr_read) begin", "            case (paddr)"]
    for reg in registers:
        lines.append(f"                OFFSET_0x{reg.offset:x}: csr_data = csr_reg_data_0x{reg.offset:x};")
    lines += ["                default: csr_data = 32'b0;", "            endcase",
              "        end", "    end", "",
              "    always_ff @(posedge pclk or negedge presetn) begin",
              "        if (!presetn)", "            prdata <= 32'b0;",
              "        else if (csr_read)", "            prdata <= csr_data;",
              "    end", ""]
    for reg in registers:
        data_name = f"csr_reg_data_0x{reg.offset:x}"
        parts = []
        for item in reg.fields:
            if item.reserved:
                continue
            bit_label = f"bit {item.lsb}" if item.width == 1 else f"bits {item.msb}:{item.lsb}"
            parts.append((item.msb, item.name, f"{bit_label}, reset {_literal(item)}, {item.access}"))
        for msb, lsb in _reserved_ranges(reg):
            width = msb - lsb + 1
            bit_label = f"bit {lsb}" if width == 1 else f"bits {msb}:{lsb}"
            parts.append((msb, f"{width}'b0", f"{bit_label}, reset {width}'h0, ro, reserved (tie0)"))
        parts.sort(key=lambda part: part[0], reverse=True)
        column_width = max(len(value) for _, value, _ in parts) + 1
        lines.append(f"    assign {data_name} = {{")
        for part_index, (_, value, comment) in enumerate(parts):
            comma = "," if part_index < len(parts) - 1 else ""
            lines.append(f"        {value + comma:<{column_width}} // {comment}")
        lines += ["    };", ""]
        for item in reg.fields:
            if item.reserved or item.access == "ro":
                continue
            state = item.name
            lines += ["    always_ff @(posedge pclk or negedge presetn) begin",
                      "        if (!presetn)", f"            {state} <= {_literal(item)};",
                      f"        else if (csr_write && paddr == OFFSET_0x{reg.offset:x})"]
            if item.access in {"rw", "rwsc"}:
                lines.append(f"            {state} <= pwdata{item.bits};")
                if item.access == "rwsc":
                    lines += ["        else", f"            {state} <= {item.width}'b0;"]
            else:
                lines += [f"            {state} <= ({state} & ~pwdata{item.bits}) | {item.name}_hwset;",
                          "        else", f"            {state} <= {state} | {item.name}_hwset;"]
            lines += ["    end", ""]
    lines += ["endmodule", ""]
    return "\n".join(lines)
