"""Read physical files into raw cells, without interpreting register fields."""
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Mapping

from openpyxl import load_workbook

from .diagnostics import TableError


@dataclass(frozen=True)
class ReadOptions:
    sheet: str = "sheet1"
    encoding: str = "utf-8"  # For future text-file adapters; XLSX ignores it.


@dataclass(frozen=True)
class RawRow:
    number: int
    cells: tuple


@dataclass(frozen=True)
class RawTable:
    path: Path
    rows: tuple[RawRow, ...]
    sheet: str | None = None


def read_xlsx_table(path, options=None):
    options = options or ReadOptions()
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        exact = [name for name in workbook.sheetnames if name == options.sheet]
        matches = exact or [name for name in workbook.sheetnames
                            if name.casefold() == options.sheet.casefold()]
        if len(matches) != 1:
            raise TableError(f"sheet {options.sheet!r} not found or ambiguous; "
                             f"available: {', '.join(workbook.sheetnames)}")
        rows = tuple(RawRow(number, tuple(cells)) for number, cells in enumerate(
            workbook[matches[0]].iter_rows(values_only=True), 1))
        return RawTable(path, rows, matches[0])
    finally:
        workbook.close()


TableReader = Callable[[Path, ReadOptions], RawTable]


def read_table(path, options=None, *, readers: Mapping[str, TableReader] | None = None):
    """Dispatch by suffix; custom adapters also return RawTable."""
    path = Path(path)
    adapters = {".xlsx": read_xlsx_table}
    if readers:
        adapters.update({suffix.lower(): adapter for suffix, adapter in readers.items()})
    adapter = adapters.get(path.suffix.lower())
    if adapter is None:
        raise TableError(f"unsupported input format {path.suffix!r}; "
                         f"supported: {', '.join(sorted(adapters))}")
    return adapter(path, options or ReadOptions())
