from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest

from csrgen.cli import main, write_template
from csrgen.generator import (Field, NumericCellWarning, Register, ReservedFieldWarning, TableError,
                              generate_rtl, read_xlsx)


@pytest.fixture
def table(tmp_path):
    path = tmp_path / "registers.xlsx"
    write_template(path)
    return path


def change_cell(path, cell, value):
    workbook = load_workbook(path)
    workbook.active[cell] = value
    workbook.save(path)
    workbook.close()


def test_screenshot_layout(table):
    regs = read_xlsx(table)
    assert [reg.offset for reg in regs] == [0, 4]
    assert [(f.name, f.width, f.reset, f.access) for f in regs[0].fields] == [
        ("a", 1, 1, "rw"), ("b", 2, 1, "ro"), ("c", 28, 0, "rw1c")]
    assert regs[1].fields[-1].reserved


def test_named_sheet_and_merged_offset(table):
    workbook = load_workbook(table)
    workbook.active.title = "Status"
    workbook.active.merge_cells("B2:B4")
    workbook.active.merge_cells("B6:B9")
    workbook.save(table)
    workbook.close()
    assert len(read_xlsx(table, sheet="status")) == 2
    with pytest.raises(TableError, match="available: Status"):
        read_xlsx(table)


@pytest.mark.parametrize("headers", [[], ["arbitrary header"] * 8])
def test_fixed_columns_ignore_headers_and_description_is_optional(tmp_path, headers):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "sheet1"
    sheet.append(headers)
    sheet.append(["ignored A column", "0x10", "data", "FF", "[7:0]", "rw", None, "ignored H column"])
    path = tmp_path / "table.xlsx"
    workbook.save(path)
    workbook.close()
    regs = read_xlsx(path)
    assert regs[0].offset == 16
    assert regs[0].fields[0].reset == 255
    assert regs[0].fields[0].row == 2
    assert regs[0].fields[0].description == ""


def test_numeric_values_are_numbers_hex_strings_are_hex(table):
    change_cell(table, "D4", 10)
    with pytest.warns(NumericCellWarning, match=r"row 4 reset: numeric cell 10 .*\(0xA\).*0x10"):
        assert read_xlsx(table)[0].fields[2].reset == 10
    change_cell(table, "D4", "10")
    assert read_xlsx(table)[0].fields[2].reset == 16


def test_numeric_offset_warns_only_when_hex_reading_differs(table, tmp_path, capsys, recwarn):
    change_cell(table, "B2", 0)
    change_cell(table, "B6", 8)
    assert [reg.offset for reg in read_xlsx(table)] == [0, 8]
    assert not recwarn.list
    change_cell(table, "B6", 20)
    with pytest.warns(NumericCellWarning, match=r"row 6 offset: numeric cell 20 .*\(0x14\).*0x20"):
        assert read_xlsx(table)[1].offset == 0x14
    assert main([str(table), "-o", str(tmp_path / "apb_csr.sv")]) == 0
    assert capsys.readouterr().err.count("csrgen: warning: row 6 offset") == 1


def test_duplicate_offset_blocks_rejected(table):
    # The same offset on consecutive rows of one block is still accepted.
    change_cell(table, "B3", "0x0")
    change_cell(table, "B4", "0")
    assert len(read_xlsx(table)[0].fields) == 3
    change_cell(table, "B9", "0x0")
    change_cell(table, "E9", "31")
    with pytest.raises(TableError, match="row 9: duplicate register offset 0x0, already defined at row 2"):
        read_xlsx(table)


@pytest.mark.parametrize("cell,value,message", [
    ("B2", None, "first field requires an offset"),
    ("B2", "0x2", "4-byte aligned"),
    ("B6", "0x10000", "exceeds aw=16"),
    ("E3", "1:0", "overlapping"),
    ("E3", "1:2", "bits must satisfy"),
    ("E3", "32", "bits must satisfy"),
    ("E3", "bad", "invalid bit range"),
    ("D2", "2", "does not fit"),
    ("D2", -1, "does not fit"),
    ("D2", 0.5, "expected an integer"),
    ("D2", "=1", "formulas are not supported"),
    ("D2", None, "reset is missing"),
    ("F2", "wo", "unsupported access"),
    ("C2", "module", "invalid Verilog identifier"),
    ("C2", "a-b", "invalid Verilog identifier"),
    ("C6", "a", "duplicate or conflicting"),
    ("C6", "c_hwset", "duplicate or conflicting"),
    ("C2", "paddr", "duplicate or conflicting"),
    ("C2", "AW", "duplicate or conflicting"),
    ("C2", "csr_read", "duplicate or conflicting"),
    ("C2", "csr_data", "duplicate or conflicting"),
    ("C2", "OFFSET_0x0", "duplicate or conflicting"),
    ("C2", "csr_reg_data_0x0", "duplicate or conflicting"),
    ("E9", "30", "overlapping"),
    ("E9", "32", "bits must satisfy"),
])
def test_bad_table_rejected(table, cell, value, message):
    change_cell(table, cell, value)
    with pytest.raises(TableError, match=message):
        read_xlsx(table)


@pytest.mark.parametrize("reset,access", [
    ("1", "ro"), ("0", "rw"), ("FF", "rw1c"), ("0", "rwsc"), ("0", "unsupported"),
    (None, None), ("invalid", "ro"), (-1, "ro"), ("=1", "ro"),
])
def test_reserved_metadata_warns_and_is_normalized(table, reset, access):
    change_cell(table, "D9", reset)
    change_cell(table, "F9", access)
    with pytest.warns(ReservedFieldWarning, match="row 9.*using reset=0, access=ro") as notices:
        regs = read_xlsx(table)
    assert len(notices) == 1
    item = regs[1].fields[-1]
    assert item.reserved and item.reset == 0 and item.access == "ro"
    # Bad metadata must not introduce a writable register or hardware-set port.
    rtl = generate_rtl(regs)
    assert "reserved_hwset" not in rtl
    assert "output reserved" not in rtl


def test_reserved_warning_cli_still_generates(table, tmp_path, capsys):
    change_cell(table, "D9", "FF")
    change_cell(table, "F9", "rw1c")
    output = tmp_path / "apb_csr.sv"
    assert main([str(table), "-o", str(output)]) == 0
    assert output.is_file()
    warnings_text = capsys.readouterr().err
    assert warnings_text.count("csrgen: warning:") == 1
    assert "row 9" in warnings_text and "tie0" in warnings_text


@pytest.mark.parametrize("access", ["rwsc", "RWSC"])
def test_rwsc_access_from_xlsx(table, tmp_path, access):
    change_cell(table, "F2", access)
    regs = read_xlsx(table)
    assert regs[0].fields[0].access == "rwsc"
    assert regs[0].fields[0].reset == 1
    assert main([str(table), "-o", str(tmp_path / "rwsc_regs.sv")]) == 0


@pytest.mark.parametrize("description", [
    "First line\nSecond line\n\nendmodule",
    "This register description contains several words and should wrap cleanly. " * 5,
    "寄存器描述较长，需要保留全部内容并按显示宽度分行。" * 5,
])
def test_description_does_not_appear_in_rtl(table, description):
    baseline = generate_rtl(read_xlsx(table))
    change_cell(table, "G2", description)
    registers = read_xlsx(table)
    assert registers[0].fields[0].description == description.strip()
    assert generate_rtl(registers) == baseline


@pytest.mark.parametrize("aw", [1, 33, 0, -1, True])
def test_invalid_aw(table, aw):
    with pytest.raises(TableError, match="aw must"):
        read_xlsx(table, aw=aw)


def test_address_width_boundary(table):
    with pytest.raises(TableError, match="exceeds aw=2"):
        read_xlsx(table, aw=2)
    assert len(read_xlsx(table, aw=3)) == 2
    change_cell(table, "B6", "FFFFFFFC")
    assert read_xlsx(table, aw=32)[1].offset == 0xFFFFFFFC


def test_empty_workbook_and_no_fields(tmp_path):
    workbook = Workbook()
    workbook.active.title = "sheet1"
    path = tmp_path / "empty.xlsx"
    workbook.save(path)
    workbook.close()
    with pytest.raises(TableError, match="no register fields"):
        read_xlsx(path)
    write_template(path)
    workbook = load_workbook(path)
    workbook.active.delete_rows(2, 20)
    workbook.save(path)
    workbook.close()
    with pytest.raises(TableError, match="no register fields"):
        read_xlsx(path)


def test_reserved_only_and_full_width(table):
    regs = [Register(0, [Field("reserved", 31, 0, 0, "ro")])]
    assert "csr_data = 32'b0;" in generate_rtl(regs)
    regs = [Register(0, [Field("data", 31, 0, 0xFFFFFFFF, "rw")])]
    assert "data <= 32'hFFFFFFFF" in generate_rtl(regs)


def test_cli_template_and_generation(tmp_path, capsys):
    table = tmp_path / "table.xlsx"
    output = tmp_path / "rtl" / "custom_regs.sv"
    assert main(["--template", str(table)]) == 0
    assert main([str(table), "--aw", "24", "--module", "custom_regs", "-o", str(output)]) == 0
    assert "parameter integer AW = 24" in output.read_text()
    assert "module custom_regs" in output.read_text()
    assert main([str(table), "--module", "module", "-o", str(output)]) == 2
    assert "invalid Verilog identifier" in capsys.readouterr().err
    original = table.read_bytes()
    assert main([str(table), "-o", str(table)]) == 2
    assert table.read_bytes() == original


def test_cli_missing_and_invalid_files(tmp_path, capsys):
    assert main([str(tmp_path / "missing.xlsx")]) == 2
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not an xlsx")
    assert main([str(invalid)]) == 2
    assert "csrgen: error:" in capsys.readouterr().err


@pytest.mark.parametrize("mode", ["script", "module", "import"])
def test_same_named_script_does_not_shadow_package(table, tmp_path, mode):
    root = Path(__file__).resolve().parents[1]
    output = tmp_path / "regs.sv"
    if mode == "import":
        args = ["-c", "from csrgen import read_xlsx; from csrgen.cli import main; "
                "raise SystemExit(main())"]
    else:
        args = ["csrgen.py"] if mode == "script" else ["-m", "csrgen"]
    result = subprocess.run(
        [sys.executable, *args, str(table), "-o", str(output)],
        cwd=root, capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stderr
    assert f"module {table.stem} #(" in output.read_text()


@pytest.mark.parametrize("options,filename,module", [
    ([], "apb_csr_common.sv", "apb_csr_common"),
    (["--module-prefix", "AAA"], "apb_csr_common.sv", "AAA_apb_csr_common"),
    (["--file-prefix", "rtl"], "rtl_apb_csr_common.sv", "apb_csr_common"),
    (["--module-prefix", "AAA", "--file-prefix", "rtl"], "rtl_apb_csr_common.sv", "AAA_apb_csr_common"),
    (["--module", "ccc"], "apb_csr_common.sv", "ccc"),
    (["--module", "ccc", "--module-prefix", "AAA"], "apb_csr_common.sv", "AAA_ccc"),
    (["--module", "ccc", "--module-prefix", "AAA", "--file-prefix", "rtl"], "rtl_apb_csr_common.sv", "AAA_ccc"),
])
def test_output_names_are_independent(table, tmp_path, monkeypatch, options, filename, module):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "apb_csr_common.xlsx"
    source.write_bytes(table.read_bytes())
    assert main([str(source), *options]) == 0
    outputs = list(tmp_path.glob("*.sv"))
    assert [output.name for output in outputs] == [filename]
    assert f"module {module} #(" in outputs[0].read_text()


@pytest.mark.parametrize("options", [
    ["--module-prefix", "bad-name"],
    ["--module-prefix", ""],
    ["--file-prefix", "../outside"],
    ["--file-prefix", "bad:name"],
    ["--file-prefix", ""],
])
def test_invalid_prefix_does_not_write_file(table, tmp_path, monkeypatch, options):
    monkeypatch.chdir(tmp_path)
    assert main([str(table), *options]) == 2
    assert not list(tmp_path.glob("*.sv"))
