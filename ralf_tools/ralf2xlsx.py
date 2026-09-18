#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ralf2xlsx.py -- Synopsys RALF (Register Abstraction Layer File) -> Excel 寄存器表

    python ralf2xlsx.py top.ralf -t <topname> -o regs.xlsx [-I incdir]...

支持的 RALF 构件 (对齐 UVM RAL Generator User Guide, Appendix A):
    field / register / regfile / memory / virtual register / block / system / domain
    属性: bits bytes access reset hard_reset soft_reset volatile enum constraint doc
          attributes endian left_to_right shared cover size initial coverpoint cross change_reset
    Tcl : # 注释, set / incr, source, [expr ...], 简单 for / if, 反斜杠续行

命令按 Tcl 语义在换行/分号处结束, 所以 "source x.ralf" / "bits 8" 这种不写分号的
写法也能正确解析。

输出 sheet:
    Summary      -- 源文件/顶层/统计, 各 block 基地址
    RegisterMap  -- 寄存器 x 字段 主表 (一个字段一行, 寄存器列纵向合并)
    Registers    -- 寄存器级汇总 (一个寄存器一行)
    Memories     -- 存储器
    Enums        -- 字段枚举值
    Warnings     -- 解析/展开过程中的所有告警
    (-s 时额外每个 block 一张 sheet)
"""

import argparse
import datetime
import html
import math
import os
import re
import sys
from collections import OrderedDict

VERSION = '2.1'

# ============================================================================
# 词法分析
# ============================================================================

SPECIAL = set('{};')


class Tok(object):
    __slots__ = ('kind', 'val', 'line', 'file')

    def __init__(self, kind, val, line, file):
        self.kind = kind
        self.val = val
        self.line = line                 # 逻辑行号 (反斜杠续行不计新行)
        self.file = file

    def __repr__(self):
        return '%s(%r)@%s:%d' % (self.kind, self.val, os.path.basename(self.file), self.line)


def tokenize(text, fname):
    """把 RALF 文本切成 token: '{' '}' ';' / word / paren

    行号是 *逻辑* 行号: 反斜杠续行不算换行, 这样上层可以按 Tcl 语义
    "一条命令到行尾结束" 来切分。
    """
    toks = []
    i, n, line = 0, len(text), 1
    while i < n:
        c = text[i]
        if c == '\\' and i + 1 < n and text[i + 1] == '\n':   # 续行
            i += 2
            continue
        if c == '\n':
            line += 1
            i += 1
            continue
        if c in ' \t\r':
            i += 1
            continue
        if c == '#':                                   # Tcl 单行注释
            j = text.find('\n', i)
            i = n if j < 0 else j
            continue
        if c in SPECIAL:
            toks.append(Tok(c, c, line, fname))
            i += 1
            continue
        if c == '"':                                   # 带引号字符串
            j, buf = i + 1, []
            while j < n and text[j] != '"':
                if text[j] == '\\' and j + 1 < n:
                    buf.append(text[j + 1])
                    j += 2
                    continue
                if text[j] == '\n':
                    line += 1
                buf.append(text[j])
                j += 1
            toks.append(Tok('word', ''.join(buf), line, fname))
            i = j + 1
            continue
        if c == '(':                                   # (hdl_path) 整体作为一个 token
            depth, j = 0, i
            while j < n:
                if text[j] == '(':
                    depth += 1
                elif text[j] == ')':
                    depth -= 1
                    if depth == 0:
                        break
                elif text[j] == '\n':
                    line += 1
                j += 1
            toks.append(Tok('paren', text[i:j + 1], line, fname))
            i = j + 1
            continue
        # 裸 word: 读到空白/特殊字符; '[' ... ']' 内部允许空白 (数组下标 / [expr ..])
        j = i
        while j < n and text[j] not in ' \t\r\n{};(#"':
            if text[j] == '\\' and j + 1 < n and text[j + 1] == '\n':
                break
            if text[j] == '[':
                depth = 0
                while j < n:
                    if text[j] == '[':
                        depth += 1
                    elif text[j] == ']':
                        depth -= 1
                        if depth == 0:
                            j += 1
                            break
                    elif text[j] == '\n':
                        line += 1
                    j += 1
                continue
            j += 1
        toks.append(Tok('word', text[i:j], line, fname))
        i = j
    return toks


# ============================================================================
# 数值 / 文本辅助
# ============================================================================

BASE = {'b': 2, 'o': 8, 'd': 10, 'h': 16}
VNUM = re.compile(r"^(?P<w>\d+)?'\s*[sS]?(?P<b>[bodhBODH])(?P<v>[0-9a-fA-FxXzZ?_]+)$")


def parse_num(s):
    """解析 RALF 数字, 返回 (value|None, width|None)

    支持 8'hFF / 'h1f / 0x10 / 0b1010 / 0o17 / 十进制 / 1Fh
    """
    if s is None:
        return None, None
    s = str(s).strip().replace('_', '')
    if not s or s.lower() == 'none':
        return None, None
    m = VNUM.match(s)
    if m:
        w = int(m.group('w')) if m.group('w') else None
        v = re.sub(r'[xXzZ?]', '0', m.group('v'))
        try:
            return int(v, BASE[m.group('b').lower()]), w
        except ValueError:
            return None, w
    for pat, b in ((r'^[+-]?0[xX][0-9a-fA-F]+$', 16), (r'^[+-]?0[bB][01]+$', 2),
                   (r'^[+-]?0[oO][0-7]+$', 8)):
        if re.match(pat, s):
            return int(s, b), None
    if re.match(r'^[+-]?\d+$', s):
        return int(s, 10), None
    if re.match(r'^[0-9a-fA-F]+[hH]$', s):
        return int(s[:-1], 16), None
    return None, None


SIZE_MUL = {'k': 1024, 'K': 1024, 'M': 1 << 20, 'm': 1 << 20, 'G': 1 << 30, 'g': 1 << 30}


def parse_size(s):
    """memory size: 64k / 2M / 1024"""
    s = str(s).strip().rstrip(';')
    mul = 1
    if s and s[-1] in SIZE_MUL:
        mul = SIZE_MUL[s[-1]]
        s = s[:-1]
    v, _ = parse_num(s)
    return None if v is None else v * mul


def hexs(v, bits=32):
    if v is None:
        return ''
    return '0x%0*X' % (max(1, (bits + 3) // 4), v)


def norm_doc(s):
    """doc {} 里是 HTML 片段, 转成一行纯文本"""
    return ' '.join(html.unescape(s or '').split())


# ============================================================================
# 数据模型
# ============================================================================

class Field(object):
    def __init__(self, name):
        self.name = name
        self.bits = None                 # 位宽 (默认 1)
        self.access = None               # rw/ro/...
        self.reset = None                # (value, raw)
        self.soft_reset = None           # (value, raw)
        self.volatile = ''
        self.enum = []                   # [(name, value)]
        self.constraints = OrderedDict()
        self.doc = ''
        self.cover = []
        self.attributes = OrderedDict()


class Register(object):
    def __init__(self, name):
        self.name = name
        self.bytes = None
        self.left_to_right = False
        self.shared = None
        self.doc = ''
        self.cover = []
        self.attributes = OrderedDict()
        self.constraints = OrderedDict()
        self.insts = []                  # field 实例 (Inst)


class RegFile(object):
    def __init__(self, name):
        self.name = name
        self.shared = None
        self.doc = ''
        self.constraints = OrderedDict()
        self.change_resets = []          # [(target_path, (value, raw))]
        self.insts = []                  # register 实例


class Memory(object):
    def __init__(self, name):
        self.name = name
        self.bits = None
        self.size = None
        self.access = 'rw'
        self.initial = None
        self.shared = None
        self.doc = ''
        self.attributes = OrderedDict()


class VReg(object):
    def __init__(self, name):
        self.name = name
        self.bytes = None
        self.left_to_right = False
        self.doc = ''
        self.insts = []


class Map(object):
    """block/system 的一个物理接口 (domain); 单接口时 name=''"""

    def __init__(self, name=''):
        self.name = name
        self.bytes = None                # 总线宽度(字节) = 一个地址单元
        self.endian = 'little'
        self.doc = ''
        self.constraints = OrderedDict()
        self.attributes = OrderedDict()
        self.cover = []
        self.change_resets = []          # [(target_path, (value, raw))]
        self.insts = []


class Block(object):
    def __init__(self, name, is_system=False):
        self.name = name
        self.is_system = is_system
        self.doc = ''
        self.attributes = OrderedDict()
        self.change_resets = []          # block-scope change_reset rules
        self.maps = OrderedDict()        # domain name -> Map

    def default_map(self):
        if not self.maps:
            self.maps[''] = Map('')
        if '' in self.maps:
            return self.maps['']
        return list(self.maps.values())[0]


class Inst(object):
    """某个构件的一次实例化"""

    def __init__(self, kind, defname):
        self.kind = kind                 # field/register/regfile/memory/vreg/block/system
        self.defname = defname
        self.name = defname              # =rename 之后的实例名
        self.arr = None                  # (lo, hi) 或 None
        self.hdl = ''
        self.offset = None               # 原始字符串 (地址单元 / bit)
        self.incr = None
        self.access_restrict = ''        # read / write
        self.mem = None                  # vreg 绑定的 memory
        self.definition = None           # 内联定义
        self.domain = ''                 # block.domain=rename
        self.line = 0
        self.file = ''


# ============================================================================
# 语法分析
# ============================================================================

class ParseError(Exception):
    pass


class Parser(object):
    def __init__(self, incdirs=None):
        self.incdirs = list(incdirs or [])
        self.toks = []
        self.pos = 0
        self.env = {}
        self.warnings = []               # [(where, msg)]
        self.cur_file = '.'
        self.src_stack = []              # 递归 source 检测
        self.files = []                  # 实际读进来的文件
        self.defs = dict((k, OrderedDict()) for k in
                         ('field', 'register', 'regfile', 'memory', 'vreg', 'block', 'system'))
        self.order = []                  # 顶层出现顺序 [(kind, name)]

    # ---------------- token 流 ----------------
    def peek(self, k=0):
        i = self.pos + k
        return self.toks[i] if i < len(self.toks) else None

    def next(self):
        t = self.peek()
        if t is None:
            raise ParseError('unexpected end of file')
        self.pos += 1
        return t

    def expect(self, kind):
        t = self.next()
        if t.kind != kind:
            raise ParseError('%s:%d: expected %r, got %r' % (t.file, t.line, kind, t.val))
        return t

    def warn(self, msg, tokn=None):
        where = ''
        if tokn is not None:
            where = '%s:%d' % (os.path.basename(tokn.file), tokn.line)
        else:
            where = os.path.basename(self.cur_file)
        self.warnings.append((where, msg))

    # ---------------- 变量 / 表达式 ----------------
    def subst(self, s):
        """$var / ${var} 替换 + [expr ...] 求值"""
        if s is None:
            return None
        out = re.sub(r'\$\{(\w+)\}', lambda m: str(self.env.get(m.group(1), m.group(0))), s)
        out = re.sub(r'\$(\w+)', lambda m: str(self.env.get(m.group(1), m.group(0))), out)
        while True:
            m = re.search(r'\[\s*expr\s+([^\[\]]*)\]', out)
            if not m:
                break
            out = out[:m.start()] + str(self.safe_expr(m.group(1))) + out[m.end():]
        return out

    def safe_expr(self, e):
        """极简表达式求值 (只用于 RALF 里的 for/if/expr)"""
        e = (e or '').strip()
        e = re.sub(r"(\d*)'[sS]?([bodhBODH])([0-9a-fA-FxXzZ_]+)",
                   lambda m: str(parse_num("'%s%s" % (m.group(2), m.group(3)))[0] or 0), e)
        e = re.sub(r'\b0[xX][0-9a-fA-F]+\b', lambda m: str(int(m.group(0), 16)), e)
        e = e.replace('&&', ' and ').replace('||', ' or ')
        e = re.sub(r'(?<![=!<>])!(?!=)', ' not ', e)
        if re.search(r'[^\w\s\.\+\-\*/%()<>=!&|\^~]', e):
            self.warn('无法求值的表达式: %s' % e)
            return 0
        try:
            return eval(e, {'__builtins__': {}}, {})
        except Exception:
            self.warn('无法求值的表达式: %s' % e)
            return 0

    # ---------------- 入口 ----------------
    def parse_file(self, path):
        path = os.path.abspath(path)
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            text = f.read()
        self.toks = tokenize(text, path)
        self.pos = 0
        self.cur_file = path
        self.src_stack.append(path)
        self.files.append(path)
        self.parse_body(None)
        self.src_stack.pop()
        return self

    def find_include(self, name, base):
        cands = [name, os.path.join(os.path.dirname(base), name)]
        cands += [os.path.join(d, name) for d in self.incdirs]
        if not name.endswith('.ralf'):
            cands += [c + '.ralf' for c in list(cands)]
        for c in cands:
            if os.path.isfile(c):
                return os.path.abspath(c)
        return None

    def run_sub(self, toks, handler, *args):
        """在独立 token 流上跑一个解析函数 (用于 source / for / if)"""
        save = (self.toks, self.pos)
        self.toks, self.pos = toks, 0
        try:
            handler(*args)
        finally:
            self.toks, self.pos = save

    # ---------------- 通用: 读 header / 块 ----------------
    def read_header(self, line):
        """读该命令的 header

        Tcl 里一条命令到行尾结束, 所以除了 '{' / ';' / '}', 换行同样终止 header
        —— 这样 "bits 8" (无分号) 和 "source x.ralf" 都不会吞掉下一条命令。
        紧跟其后单独一行的 '{' 仍然算这条命令的 body。
        """
        hdr = []
        while True:
            t = self.peek()
            if t is None:
                return hdr, False
            if t.kind == '{':
                self.next()
                return hdr, True
            if t.kind == ';':
                self.next()
                return hdr, False
            if t.kind == '}':
                return hdr, False
            if t.line != line:
                return hdr, False
            hdr.append(self.next())

    def skip_block_raw(self):
        """吃掉一个 {...} 并返回内部原文"""
        depth, out = 1, []
        while depth:
            t = self.next()
            if t.kind == '{':
                depth += 1
                out.append('{')
            elif t.kind == '}':
                depth -= 1
                if depth:
                    out.append('}')
            else:
                out.append(t.val)
        return ' '.join(out).strip()

    def grab_brace_tokens(self):
        """已消费 '{', 收集到配对 '}' 的 token 列表 (不含最外层花括号)"""
        depth, out = 1, []
        while depth:
            t = self.next()
            if t.kind == '{':
                depth += 1
            elif t.kind == '}':
                depth -= 1
                if depth == 0:
                    break
            out.append(t)
        return out

    # ---------------- header 解码 ----------------
    def decode_header(self, hdr):
        """name[=rename][[n]|[m:n]] [(hdl)] [@off[+incr]] [+incr] [read|write] [mem@off]"""
        d = dict(name=None, rename=None, arr=None, hdl='', offset=None, incr=None,
                 access='', mem=None, extra=[], domain='')
        if not hdr:
            return d
        first = self.subst(hdr[0].val)
        m = re.match(r'^([^\[\]=@]+)(?:=([^\[\]@]+))?(?:\[([^\]]*)\])?$', first)
        if m:
            d['name'] = m.group(1)
            d['rename'] = m.group(2)
            if m.group(3):
                d['arr'] = self.decode_range(m.group(3))
        else:
            d['name'] = first
        if d['name'] and '.' in d['name']:
            base, dom = d['name'].split('.', 1)
            d['name'], d['domain'] = base, dom
        for t in hdr[1:]:
            v = self.subst(t.val)
            if t.kind == 'paren':
                d['hdl'] = v[1:-1].strip()
                continue
            if v.startswith('@'):
                body = v[1:]
                if '+' in body:
                    body, inc = body.split('+', 1)
                    d['incr'] = inc
                d['offset'] = body
                continue
            if v.startswith('+'):
                d['incr'] = v[1:]
                continue
            if v in ('read', 'write'):
                d['access'] = v
                continue
            if '@' in v:                        # vreg: mem@offset[+incr]
                mem, off = v.split('@', 1)
                if '+' in off:
                    off, inc = off.split('+', 1)
                    d['incr'] = inc
                mm = re.match(r'^([^\[]+)\[([^\]]*)\]$', mem)
                d['mem'] = mm.group(1) if mm else mem
                d['offset'] = off
                continue
            d['extra'].append(v)
        return d

    def decode_range(self, s):
        s = s.strip()
        if ':' in s:
            a, b = s.split(':', 1)
            lo, _ = parse_num(self.subst(a))
            hi, _ = parse_num(self.subst(b))
            if lo is None or hi is None:
                return None
            return (min(lo, hi), max(lo, hi))
        v, _ = parse_num(self.subst(s))
        if v is None:
            return None
        return (0, v - 1)

    def num(self, s):
        return parse_num(self.subst(s))[0]

    # ---------------- 主体 ----------------
    def parse_body(self, ctx):
        """ctx: None=全局 / Map / Register / RegFile / VReg / Block"""
        while True:
            t = self.peek()
            if t is None or t.kind == '}':
                return
            if t.kind == ';':
                self.next()
                continue
            if t.kind != 'word':
                raise ParseError('%s:%d: unexpected %r' % (t.file, t.line, t.val))
            kw_tok = self.next()
            kw = kw_tok.val

            if kw == 'source':
                self.do_source(ctx, kw_tok)
            elif kw == 'set':
                name = self.next().val
                self.env[name] = self.subst(self.next().val)
                if self.peek() and self.peek().kind == ';':
                    self.next()
            elif kw == 'incr':
                name = self.next().val
                step = 1
                nt = self.peek()
                if nt and nt.kind == 'word' and nt.line == kw_tok.line and \
                        re.match(r'^-?\d+$', nt.val):
                    step = int(self.next().val)
                self.env[name] = str(int(self.env.get(name, 0)) + step)
            elif kw == 'for':
                self.do_for(ctx, kw_tok)
            elif kw == 'if':
                self.do_if(ctx, kw_tok)
            elif kw in ('puts', 'unset'):
                self.read_header(kw_tok.line)
            elif kw == 'virtual':
                nxt = self.next().val
                if nxt != 'register':
                    raise ParseError('expected "virtual register", got "virtual %s"' % nxt)
                self.do_construct('vreg', ctx, kw_tok)
            elif kw in ('field', 'register', 'regfile', 'memory', 'block', 'system'):
                self.do_construct(kw, ctx, kw_tok)
            elif kw == 'domain':
                self.do_domain(ctx, kw_tok)
            else:
                self.do_property(kw, ctx, kw_tok)

    def do_source(self, ctx, kw_tok):
        t = self.peek()
        if t is None or t.kind != 'word' or t.line != kw_tok.line:
            self.warn('source 缺少文件名', kw_tok)
            return
        raw = self.subst(self.next().val)
        if self.peek() and self.peek().kind == ';':
            self.next()
        path = self.find_include(raw, self.cur_file)
        if not path:
            self.warn('source 找不到文件: %s' % raw, kw_tok)
            return
        if path in self.src_stack:
            self.warn('检测到递归 source, 已跳过: %s' % raw, kw_tok)
            return
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            sub = tokenize(f.read(), path)
        prev = self.cur_file
        self.cur_file = path
        self.src_stack.append(path)
        self.files.append(path)
        self.run_sub(sub, self.parse_body, ctx)
        self.src_stack.pop()
        self.cur_file = prev

    def do_for(self, ctx, kw_tok):
        groups = []
        for _ in range(4):
            while self.peek() and self.peek().kind == ';':
                self.next()
            if not self.peek() or self.peek().kind != '{':
                self.warn('无法解析的 for 循环, 已跳过', kw_tok)
                return
            self.next()
            groups.append(self.grab_brace_tokens())
        init, cond, step, body = groups
        try:
            var = init[1].val
            start = int(self.num(init[2].val))
            op = cond[1].val
            limit = int(self.num(self.subst(' '.join(x.val for x in cond[2:]))))
            inc = int(self.num(step[2].val)) if len(step) >= 3 else 1
        except Exception:
            self.warn('无法解析的 for 循环, 已跳过', kw_tok)
            return
        cmp_ = {'<': lambda a, b: a < b, '<=': lambda a, b: a <= b,
                '>': lambda a, b: a > b, '>=': lambda a, b: a >= b,
                '!=': lambda a, b: a != b}.get(op)
        if cmp_ is None or inc == 0:
            self.warn('不支持的 for 条件: %s' % op, kw_tok)
            return
        i, guard = start, 0
        while cmp_(i, limit) and guard < 100000:
            self.env[var] = str(i)
            self.run_sub(list(body), self.parse_body, ctx)
            i += inc
            guard += 1

    def do_if(self, ctx, kw_tok):
        if not self.peek() or self.peek().kind != '{':
            self.warn('无法解析的 if, 已跳过', kw_tok)
            return
        self.next()
        cond = self.grab_brace_tokens()
        if not self.peek() or self.peek().kind != '{':
            self.warn('无法解析的 if body, 已跳过', kw_tok)
            return
        self.next()
        body = self.grab_brace_tokens()
        else_body = None
        nt = self.peek()
        if nt and nt.kind == 'word' and nt.val == 'else':
            self.next()
            if self.peek() and self.peek().kind == '{':
                self.next()
                else_body = self.grab_brace_tokens()
        expr = self.subst(' '.join(t.val for t in cond))
        if bool(self.safe_expr(expr)):
            self.run_sub(list(body), self.parse_body, ctx)
        elif else_body is not None:
            self.run_sub(list(else_body), self.parse_body, ctx)

    def do_domain(self, ctx, kw_tok):
        hdr, opened = self.read_header(kw_tok.line)
        d = self.decode_header(hdr)
        if not isinstance(ctx, Block):
            raise ParseError('%s:%d: domain 只能出现在 block/system 内' %
                             (kw_tok.file, kw_tok.line))
        mp = Map(d['name'] or '')
        if '' in ctx.maps and not ctx.maps[''].insts and ctx.maps[''].bytes is None:
            del ctx.maps['']                     # 丢掉占位的默认 map
        ctx.maps[mp.name] = mp
        if opened:
            self.parse_body(mp)
            self.expect('}')

    # ---------------- 构件 ----------------
    def do_construct(self, kind, ctx, kw_tok):
        hdr, opened = self.read_header(kw_tok.line)
        d = self.decode_header(hdr)
        name = d['name']
        obj = None
        if opened:
            obj = self.new_def(kind, name)
            self.parse_body(obj)
            self.expect('}')

        if ctx is None:                          # 顶层定义
            if obj is None:
                if kind != 'field':              # "field tx_en;" = 1bit rw, 合法
                    self.warn('顶层 %s %s 缺少定义体, 已忽略' % (kind, name), kw_tok)
                    return
                obj = self.new_def(kind, name)
            self.defs[kind][name] = obj
            self.order.append((kind, name))
            return

        inst = Inst(kind, name)
        inst.name = d['rename'] or name
        inst.arr = d['arr']
        inst.hdl = d['hdl']
        inst.offset = d['offset']
        inst.incr = d['incr']
        inst.access_restrict = d['access']
        inst.mem = d['mem']
        inst.domain = d['domain']
        inst.definition = obj
        inst.line = kw_tok.line
        inst.file = kw_tok.file
        if obj is not None and kind in ('block', 'system'):
            self.defs[kind][name] = obj          # 内联子 block 也登记, 便于复用
        target = ctx.default_map() if isinstance(ctx, Block) else ctx
        if hasattr(target, 'insts'):
            target.insts.append(inst)
        else:
            self.warn('%s %s 出现在不支持的上下文中' % (kind, name), kw_tok)

    def new_def(self, kind, name):
        if kind == 'field':
            return Field(name)
        if kind == 'register':
            return Register(name)
        if kind == 'regfile':
            return RegFile(name)
        if kind == 'memory':
            return Memory(name)
        if kind == 'vreg':
            return VReg(name)
        if kind in ('block', 'system'):
            b = Block(name, is_system=(kind == 'system'))
            b.maps[''] = Map('')
            return b
        raise ParseError('unknown kind %s' % kind)

    # ---------------- 属性 ----------------
    def do_property(self, kw, ctx, kw_tok):
        hdr, opened = self.read_header(kw_tok.line)
        vals = [self.subst(t.val) for t in hdr]
        v0 = vals[0] if vals else ''

        # 单接口 block/system 的 bytes/endian 属于其默认 physical interface。
        # domain 内 ctx 本身就是 Map。
        prop_ctx = ctx.default_map() if isinstance(ctx, Block) and kw in ('bytes', 'endian') else ctx

        if kw == 'change_reset':
            if opened:
                self.skip_block_raw()
            if len(vals) < 2:
                self.warn('change_reset 缺少 target/value', kw_tok)
                return
            target, raw = vals[0], vals[1]
            rule = (target, (parse_num(raw)[0], raw))
            if hasattr(ctx, 'change_resets'):
                ctx.change_resets.append(rule)
            else:
                self.warn('change_reset 出现在不支持的上下文中', kw_tok)
            return

        if kw == 'constraint':
            body = self.skip_block_raw() if opened else ''
            if hasattr(ctx, 'constraints'):
                ctx.constraints[v0 or 'c%d' % len(ctx.constraints)] = body
            return
        if kw == 'doc':
            body = self.skip_block_raw() if opened else ' '.join(vals)
            if hasattr(ctx, 'doc'):
                ctx.doc = (ctx.doc + ' ' + norm_doc(body)).strip()
            return
        if kw in ('coverpoint', 'cross'):
            if opened:
                self.skip_block_raw()
            return
        if kw == 'attributes':
            body = self.skip_block_raw() if opened else ''
            if hasattr(ctx, 'attributes'):
                for item in body.split(','):
                    parts = item.split()
                    if parts:
                        ctx.attributes[parts[0]] = ' '.join(parts[1:])
            return
        if kw == 'enum':
            body = self.skip_block_raw() if opened else ' '.join(vals)
            self.parse_enum(body, ctx, kw_tok)
            return

        if opened:                                # 未知的带块属性
            self.skip_block_raw()
            self.warn('忽略未知属性块: %s' % kw, kw_tok)
            return

        if kw == 'bits':
            ctx.bits = self.num(v0)
        elif kw == 'bytes':
            prop_ctx.bytes = self.num(v0)
        elif kw == 'size':
            ctx.size = parse_size(v0)
        elif kw == 'access':
            ctx.access = v0.lower()
        elif kw in ('reset', 'hard_reset'):
            ctx.reset = (parse_num(v0)[0], v0)
        elif kw == 'soft_reset':
            ctx.soft_reset = (parse_num(v0)[0], v0)
        elif kw == 'volatile':
            ctx.volatile = v0 or 'yes'
        elif kw == 'initial':
            ctx.initial = v0
        elif kw == 'left_to_right':
            ctx.left_to_right = True
        elif kw == 'endian':
            prop_ctx.endian = v0
        elif kw == 'shared':
            ctx.shared = v0 or 'yes'
        elif kw == 'cover':
            if hasattr(ctx, 'cover'):
                ctx.cover.extend(vals)
        elif kw in ('noise', 'hdl_path', 'individually_accessible'):
            pass
        else:
            self.warn('忽略未知属性: %s %s' % (kw, ' '.join(vals)), kw_tok)

    def parse_enum(self, body, ctx, kw_tok=None):
        """enum { AA, BB=2, CC } / enum { OFF = 0, ON = 1 }  (= 两边可能有空格)"""
        if not hasattr(ctx, 'enum'):
            return
        nxt = 0
        for item in (body or '').split(','):
            item = item.strip().strip(';').strip()
            if not item:
                continue
            m = re.match(r'^([A-Za-z_]\w*)\s*(?:=\s*(.+))?$', item)
            if not m:
                self.warn('无法解析的 enum 项: %s' % item, kw_tok)
                continue
            val = nxt
            if m.group(2) is not None:
                v = parse_num(self.subst(m.group(2).strip()))[0]
                if v is not None:
                    val = v
            ctx.enum.append((m.group(1), val))
            nxt = val + 1


# ============================================================================
# 展开 (elaboration)
# ============================================================================

class Elaborator(object):
    """把定义 + 实例展开成扁平的寄存器/字段表

    addr_mode:
      'byte' -- RALF 里的 @offset 直接当字节偏移 (手写 / IP-XACT 转来的 RALF 最常见)
      'word' -- @offset 是地址单元(= 所在 block 的 bytes)个数, 即 ralgen 不加 -B 时的语义
    expand_arrays:
      False 时数组只出一行, 名字写成 name[lo:hi], 地址取第 lo 个元素
    """

    def __init__(self, parser, addr_mode='byte', expand_arrays=True):
        self.p = parser
        self.addr_mode = addr_mode
        self.expand = expand_arrays
        self.rows = []                   # 字段级行
        self.reg_rows = []               # 寄存器级行
        self.mem_rows = []
        self.blk_rows = []
        self.warnings = []
        self._seen = set()
        self.reset_rules = []             # [(full_target, value, raw)]

    def warn(self, m):
        if m not in self.warnings:
            self.warnings.append(m)

    # ---------- 地址单位 ----------
    def ubytes(self, bus):
        """一个 "地址单元" 对应多少字节"""
        return 1 if self.addr_mode == 'byte' else bus

    def stride(self, width_bits, bus):
        """一个 width_bits 宽的寄存器占用多少地址单元 (总线对齐)"""
        words = max(1, int(math.ceil((width_bits / 8.0) / bus)))
        return words * bus if self.addr_mode == 'byte' else words

    # ---------- 数组 ----------
    def idxs(self, arr):
        if arr is None:
            return [None]
        return list(range(arr[0], arr[1] + 1)) if self.expand else [arr[0]]

    def anm(self, name, arr, k):
        if arr is None:
            return name
        if not self.expand:
            return '%s[%d:%d]' % (name, arr[0], arr[1])
        return '%s[%d]' % (name, k)

    # ---------- field 布局 ----------
    def layout(self, reg):
        """返回 ([(inst_name, Field, lsb, bits)], 寄存器位宽)"""
        out = []
        items = []
        for inst in reg.insts:
            if inst.kind != 'field':
                continue
            fd = inst.definition or self.p.defs['field'].get(inst.defname)
            if fd is None:
                self.warn('找不到 field 定义: %s (in register %s)' % (inst.defname, reg.name))
                fd = Field(inst.defname)
            items.append((inst, fd))

        width = reg.bytes * 8 if reg.bytes else None

        if not reg.left_to_right:
            nxt = 0
            for inst, fd in items:
                bits = fd.bits or 1
                base = self.p.num(inst.offset) if inst.offset not in (None, 'none', '') else None
                if base is None:
                    base = nxt
                incr = (self.p.num(inst.incr) or bits) if inst.incr else bits
                lo = inst.arr[0] if inst.arr else 0
                last = base
                for k in self.idxs(inst.arr):
                    lsb = base + ((k - lo) if k is not None else 0) * incr
                    out.append((self.anm(inst.name, inst.arr, k), fd, lsb, bits))
                    last = lsb
                if inst.arr and not self.expand:
                    last = base + (inst.arr[1] - inst.arr[0]) * incr
                nxt = last + bits
        else:
            # left_to_right: offset 是“从最高有效位向右”的距离。先在这个
            # top-relative 坐标系中完成布局（包含显式 gap），再整体向 LSB 对齐。
            # 不能简单用 sum(field bits) 作为 total，否则 @offset 产生的 gap 会丢失。
            top_items = []                       # (name, fd, top_off, bits)
            nxt_top = 0
            span = 0
            for inst, fd in items:
                bits = fd.bits or 1
                off = (self.p.num(inst.offset)
                       if inst.offset not in (None, 'none', '') else None)
                base = nxt_top if off is None else off
                incr = (self.p.num(inst.incr) or bits) if inst.incr else bits
                lo = inst.arr[0] if inst.arr else 0
                cnt = (inst.arr[1] - inst.arr[0] + 1) if inst.arr else 1

                for k in self.idxs(inst.arr):
                    j = (k - lo) if k is not None else 0
                    top_off = base + j * incr
                    top_items.append((self.anm(inst.name, inst.arr, k), fd, top_off, bits))

                # 即使 --no-expand-arrays，也必须把未显示元素算进实际跨度。
                last_top = base + (cnt - 1) * incr
                nxt_top = last_top + bits
                span = max(span, nxt_top)

            for nm, fd, top_off, bits in top_items:
                lsb = span - (top_off + bits)
                out.append((nm, fd, max(lsb, 0), bits))

        if width is None:
            msb = max([l + b for _, _, l, b in out] or [8])
            width = int(math.ceil(msb / 8.0)) * 8
        out.sort(key=lambda x: x[2])
        return out, width

    def _reg_def(self, inst):
        return inst.definition or self.p.defs['register'].get(inst.defname)

    def add_reset_rules(self, base_path, rules):
        for target, (value, raw) in (rules or []):
            if not self.expand and re.search(r'\[\d+\]', target):
                self.warn('数组未展开时无法精确表示 change_reset %s; 保留原 reset' % target)
            full = '%s.%s' % (base_path, target) if base_path else target
            self.reset_rules.append((full, value, raw))

    def reset_override(self, reg_path, field_path, lsb, bits):
        """Return (value, raw) for the most specific/latest matching change_reset.

        A register-level rule is sliced down to the field bits; a field-level rule
        directly supplies the field reset value.
        """
        for target, value, raw in reversed(self.reset_rules):
            if target == field_path:
                return value, raw
            if target == reg_path:
                if value is None:
                    return None, raw
                mask = (1 << bits) - 1
                return (value >> lsb) & mask, raw
        return None

    # ---------- 层次遍历 ----------
    def walk_block(self, blk, path, base_byte, parent_bytes, domain=''):
        key = (id(blk), path, domain)
        if key in self._seen:
            self.warn('检测到重复/递归实例化, 已跳过: %s' % path)
            return
        self._seen.add(key)

        # block-scope change_reset 以该 block 实例路径为基准。
        self.add_reset_rules(path, getattr(blk, 'change_resets', []))

        if domain:
            if domain not in blk.maps:
                self.warn('找不到 domain %s (in %s)' % (domain, path))
                return
            maps = [(domain, blk.maps[domain])]
        else:
            maps = list(blk.maps.items())

        for mname, mp in maps:
            bus = mp.bytes or parent_bytes or 4
            bpath = path if not mname else '%s.%s' % (path, mname)
            if bpath != path:
                self.add_reset_rules(bpath, getattr(blk, 'change_resets', []))
            self.add_reset_rules(bpath, getattr(mp, 'change_resets', []))
            self.blk_rows.append(dict(block=bpath, kind='system' if blk.is_system else 'block',
                                      bus_bytes=bus, endian=mp.endian, base=base_byte,
                                      items=len(mp.insts), doc=blk.doc))
            self.walk_map(mp, bpath, base_byte, bus, '')

    def walk_map(self, mp, bpath, base_byte, bus, prefix):
        nxt_unit = 0                                  # 下一个可用地址单元
        for inst in mp.insts:
            off = inst.offset
            explicit = off not in (None, 'none', '')
            unit = self.p.num(off) if explicit else nxt_unit
            if off == 'none':
                unit = None
            elif explicit and unit is None:
                self.warn('无法解析的地址偏移 @%s (%s.%s)' % (off, bpath, inst.name))
                unit = nxt_unit

            if inst.kind == 'register':
                nxt_unit = self.emit_register(inst, bpath, base_byte, bus, prefix,
                                              unit, nxt_unit, mp.endian)
            elif inst.kind == 'regfile':
                nxt_unit = self.emit_regfile(inst, bpath, base_byte, bus, prefix,
                                             unit, nxt_unit, mp.endian)
            elif inst.kind == 'memory':
                nxt_unit = self.emit_memory(inst, bpath, base_byte, bus, prefix,
                                            unit, nxt_unit)
            elif inst.kind == 'vreg':
                self.emit_vreg(inst, bpath, base_byte, bus, prefix, mp.endian)
            elif inst.kind in ('block', 'system'):
                sub = (inst.definition or self.p.defs[inst.kind].get(inst.defname)
                       or self.p.defs['block'].get(inst.defname)
                       or self.p.defs['system'].get(inst.defname))
                if sub is None:
                    self.warn('找不到 %s 定义: %s' % (inst.kind, inst.defname))
                    continue
                incr = self.p.num(inst.incr) if inst.incr else 0
                lo = inst.arr[0] if inst.arr else 0
                for k in self.idxs(inst.arr):
                    j = (k - lo) if k is not None else 0
                    b = base_byte + ((unit or 0) + j * (incr or 0)) * self.ubytes(bus)
                    self.walk_block(sub, '%s.%s' % (bpath, self.anm(inst.name, inst.arr, k)),
                                    b, bus, inst.domain)

    # ---------- register ----------
    def emit_register(self, inst, bpath, base_byte, bus, prefix, unit, nxt_unit, endian):
        rd = self._reg_def(inst)
        if rd is None:
            self.warn('找不到 register 定义: %s' % inst.defname)
            return nxt_unit
        fields, width = self.layout(rd)
        naddr = self.stride(width, bus)
        incr = (self.p.num(inst.incr) or naddr) if inst.incr else naddr
        lo = inst.arr[0] if inst.arr else 0
        last = unit
        for k in self.idxs(inst.arr):
            j = (k - lo) if k is not None else 0
            u = None if unit is None else unit + j * incr
            last = u
            self.emit_reg_rows(rd, self.anm(inst.name, inst.arr, k), fields, width,
                               bpath, prefix, base_byte, bus, u, inst, endian)
        if unit is None:
            return nxt_unit
        if inst.arr and not self.expand:
            last = unit + (inst.arr[1] - inst.arr[0]) * incr
        return (last if last is not None else unit) + naddr

    def emit_reg_rows(self, rd, nm, fields, width, bpath, prefix, base_byte, bus,
                      u, inst, endian):
        byte_off = None if u is None else u * self.ubytes(bus)
        absaddr = None if u is None else base_byte + byte_off
        rdoc = (rd.doc + (' [shared]' if getattr(rd, 'shared', None) else '')).strip()
        common = dict(block=bpath, reg=prefix + nm, regtype=rd.name,
                      byte_off='' if byte_off is None else hexs(byte_off),
                      addr='' if absaddr is None else hexs(absaddr),
                      width=width, hdl=inst.hdl,
                      restrict=inst.access_restrict, rdoc=rdoc,
                      _addr=absaddr if absaddr is not None else -1)
        self.reg_rows.append(dict(common, bytes=width // 8, endian=endian,
                                  nfields=len(fields),
                                  attrs=', '.join('%s=%s' % kv for kv in rd.attributes.items()),
                                  where='%s:%d' % (os.path.basename(inst.file), inst.line)
                                        if inst.file else ''))
        if not fields:
            r = dict(common)
            r.update(field='-', msb='', lsb='', bits='', fwidth='', access='',
                     reset='', reset_raw='', soft_reset='', volatile='', enum='',
                     constr='', fdoc='')
            self.rows.append(r)
            return
        reg_path = '%s.%s%s' % (bpath, prefix, nm)
        for fname, fd, lsb, bits in fields:
            msb = lsb + bits - 1
            acc = (fd.access or ('ro' if fd.name in ('unused', 'reserved') else 'rw')).upper()
            field_path = '%s.%s' % (reg_path, fname)
            ov = self.reset_override(reg_path, field_path, lsb, bits)
            if ov is not None:
                ov_val, ov_raw = ov
                rst = hexs(ov_val, bits) if ov_val is not None else ov_raw
                rst_raw = ov_raw
            elif fd.reset and fd.reset[0] is not None:
                rst, rst_raw = hexs(fd.reset[0], bits), fd.reset[1]
            elif fd.reset:
                rst, rst_raw = fd.reset[1], fd.reset[1]
            else:
                rst, rst_raw = hexs(0, bits), ''
            if fd.soft_reset and fd.soft_reset[0] is not None:
                srst = hexs(fd.soft_reset[0], bits)
            else:
                srst = fd.soft_reset[1] if fd.soft_reset else ''
            r = dict(common)
            r.update(field=fname, msb=msb, lsb=lsb,
                     bits='[%d:%d]' % (msb, lsb) if bits > 1 else '[%d]' % lsb,
                     fwidth=bits, access=acc, reset=rst, reset_raw=rst_raw,
                     soft_reset=srst, volatile=fd.volatile,
                     enum=', '.join('%s=%s' % (n, hexs(v, bits)) for n, v in fd.enum),
                     constr='; '.join(fd.constraints.keys()),
                     fdoc=fd.doc)
            self.rows.append(r)

    # ---------- regfile ----------
    def emit_regfile(self, inst, bpath, base_byte, bus, prefix, unit, nxt_unit, endian):
        rf = inst.definition or self.p.defs['regfile'].get(inst.defname)
        if rf is None:
            self.warn('找不到 regfile 定义: %s' % inst.defname)
            return nxt_unit
        span = self.rf_span(rf, bus)
        incr = (self.p.num(inst.incr) or span) if inst.incr else span
        lo = inst.arr[0] if inst.arr else 0
        last_end = nxt_unit
        for k in self.idxs(inst.arr):
            j = (k - lo) if k is not None else 0
            nm = self.anm(inst.name, inst.arr, k)
            u = None if unit is None else unit + j * incr
            # regfile-scope change_reset 相对该 regfile 实例路径。
            rf_base_path = '%s.%s%s' % (bpath, prefix, nm)
            self.add_reset_rules(rf_base_path, getattr(rf, 'change_resets', []))
            sub_next = 0
            for ri in rf.insts:
                if ri.kind != 'register':
                    continue
                rexp = ri.offset not in (None, 'none', '')
                runit = self.p.num(ri.offset) if rexp else sub_next
                if ri.offset == 'none':
                    runit = None
                elif runit is None:
                    runit = sub_next
                abs_unit = None if (u is None or runit is None) else u + runit
                rd = self._reg_def(ri)
                if rd is None:
                    self.warn('找不到 register 定义: %s' % ri.defname)
                    continue
                fields, width = self.layout(rd)
                naddr = self.stride(width, bus)
                rincr = (self.p.num(ri.incr) or naddr) if ri.incr else naddr
                rlo = ri.arr[0] if ri.arr else 0
                for m in self.idxs(ri.arr):
                    jj = (m - rlo) if m is not None else 0
                    uu = None if abs_unit is None else abs_unit + jj * rincr
                    self.emit_reg_rows(rd, self.anm(ri.name, ri.arr, m), fields, width,
                                       bpath, prefix + nm + '.', base_byte, bus, uu,
                                       ri, endian)
                if runit is not None:
                    cnt = (ri.arr[1] - ri.arr[0]) if ri.arr else 0
                    sub_next = runit + cnt * rincr + naddr
            if u is not None:
                last_end = u + span
        if unit is None:
            return nxt_unit
        if inst.arr and not self.expand:
            last_end = unit + (inst.arr[1] - inst.arr[0]) * incr + span
        return last_end

    def rf_span(self, rf, bus):
        """regfile 占用的地址单元数"""
        nxt = 0
        for ri in rf.insts:
            if ri.kind != 'register':
                continue
            rd = self._reg_def(ri)
            if rd is None:
                continue
            _, width = self.layout(rd)
            naddr = self.stride(width, bus)
            base = self.p.num(ri.offset) if ri.offset not in (None, 'none', '') else nxt
            if base is None:
                base = nxt
            cnt = (ri.arr[1] - ri.arr[0]) if ri.arr else 0
            incr = (self.p.num(ri.incr) or naddr) if ri.incr else naddr
            nxt = base + cnt * incr + naddr
        return max(1, nxt)

    # ---------- memory / vreg ----------
    def emit_memory(self, inst, bpath, base_byte, bus, prefix, unit, nxt_unit):
        md = inst.definition or self.p.defs['memory'].get(inst.defname)
        if md is None:
            self.warn('找不到 memory 定义: %s' % inst.defname)
            return nxt_unit
        loc_units = self.stride(md.bits or 8, bus)
        naddr = (md.size or 0) * loc_units
        byte_off = None if unit is None else unit * self.ubytes(bus)
        total = ''
        if md.size and md.bits:
            nbytes = md.size * md.bits / 8.0
            total = '%.3g KB' % (nbytes / 1024.0) if nbytes >= 1024 else '%d B' % nbytes
        self.mem_rows.append(dict(
            block=bpath, name=prefix + inst.name, memtype=md.name,
            addr='' if unit is None else hexs(base_byte + byte_off),
            byte_off='' if unit is None else hexs(byte_off),
            size=md.size, bits=md.bits, total=total,
            access=(inst.access_restrict or md.access or 'rw').upper(),
            initial=md.initial or '', hdl=inst.hdl, doc=md.doc))
        if unit is None:
            return nxt_unit
        return unit + naddr

    def emit_vreg(self, inst, bpath, base_byte, bus, prefix, endian):
        vd = inst.definition or self.p.defs['vreg'].get(inst.defname)
        if vd is None:
            self.warn('找不到 virtual register 定义: %s' % inst.defname)
            return
        shim = Register(vd.name)
        shim.insts = vd.insts
        shim.bytes = vd.bytes
        shim.left_to_right = vd.left_to_right
        shim.doc = ('virtual register in memory "%s"%s' %
                    (inst.mem or '?', (' - ' + vd.doc) if vd.doc else ''))
        fields, width = self.layout(shim)
        for k in self.idxs(inst.arr):
            self.emit_reg_rows(shim, self.anm(inst.name, inst.arr, k), fields, width,
                               bpath, prefix, base_byte, bus, None, inst, endian)


# ============================================================================
# 表格定义
# ============================================================================

REG_COLS = [
    ('block',      'Block',          22),
    ('reg',        'Register',       30),
    ('addr',       'Address',        14),
    ('byte_off',   'Offset',         12),
    ('width',      'Width(bit)',      9),
    ('field',      'Field',          22),
    ('bits',       'Bits',           10),
    ('msb',        'MSB',             6),
    ('lsb',        'LSB',             6),
    ('fwidth',     'Size',            6),
    ('access',     'Access',          9),
    ('reset',      'Reset',          12),
    ('reset_raw',  'Reset(RALF)',    13),
    ('soft_reset', 'Soft Reset',     12),
    ('volatile',   'Volatile',        9),
    ('enum',       'Enum',           26),
    ('constr',     'Constraint',     14),
    ('restrict',   'R/W Limit',      10),
    ('hdl',        'HDL Path',       26),
    ('fdoc',       'Field Desc',     34),
    ('rdoc',       'Register Desc',  30),
]
MERGE_KEYS = ('block', 'reg', 'addr', 'byte_off', 'width', 'restrict', 'hdl', 'rdoc')
CENTER_KEYS = ('addr', 'byte_off', 'width', 'bits', 'msb', 'lsb', 'fwidth', 'access',
               'reset', 'reset_raw', 'soft_reset', 'volatile', 'size', 'total', 'value',
               'bytes', 'nfields', 'endian')
MONO_KEYS = ('addr', 'byte_off', 'reset', 'reset_raw', 'bits', 'value')

REGSUM_COLS = [
    ('block', 'Block', 22), ('reg', 'Register', 30), ('regtype', 'Type', 20),
    ('addr', 'Address', 14), ('byte_off', 'Offset', 12), ('bytes', 'Bytes', 7),
    ('width', 'Width(bit)', 10), ('nfields', 'Fields', 7), ('endian', 'Endian', 9),
    ('restrict', 'R/W Limit', 10), ('attrs', 'Attributes', 20),
    ('hdl', 'HDL Path', 26), ('rdoc', 'Description', 40), ('where', 'Source', 18),
]

MEM_COLS = [
    ('block', 'Block', 22), ('name', 'Memory', 24), ('memtype', 'Type', 18),
    ('addr', 'Address', 14), ('byte_off', 'Offset', 12), ('size', 'Size(locs)', 12),
    ('bits', 'Width(bit)', 10), ('total', 'Total', 12), ('access', 'Access', 9),
    ('initial', 'Initial', 10), ('hdl', 'HDL Path', 26), ('doc', 'Description', 30),
]

ENUM_COLS = [('block', 'Block', 22), ('reg', 'Register', 28), ('field', 'Field', 22),
             ('name', 'Enum Name', 24), ('value', 'Value', 12)]

WARN_COLS = [('kind', 'Stage', 10), ('where', 'Where', 26), ('msg', 'Warning', 100)]


def bit_range_of(r):
    return r.get('bits', '')


# ============================================================================
# Excel 输出
# ============================================================================

def write_xlsx(path, ela, meta, warn_rows, sheet_per_block=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    hdr_font = Font(bold=True, color='FFFFFF', size=10)
    hdr_fill = PatternFill('solid', fgColor='2F5597')
    reg_fill = PatternFill('solid', fgColor='DCE6F1')
    warn_fill = PatternFill('solid', fgColor='FFF2CC')
    thin = Side(style='thin', color='B7B7B7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    mono = Font(name='Consolas', size=10)

    wb = Workbook()

    def dump(ws, rows, cols, merge=False, fill=None):
        for j, (_, h, w) in enumerate(cols, 1):
            c = ws.cell(1, j, h)
            c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, border
            ws.column_dimensions[get_column_letter(j)].width = w
        for i, row in enumerate(rows, start=2):
            for j, (key, _, _) in enumerate(cols, 1):
                c = ws.cell(i, j, row.get(key, ''))
                c.border = border
                c.alignment = center if key in CENTER_KEYS else left
                if key in MONO_KEYS:
                    c.font = mono
                if fill is not None:
                    c.fill = fill
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 22
        if rows:
            ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), len(rows) + 1)
        if not merge or not rows:
            return
        idx = dict((k, j) for j, (k, _, _) in enumerate(cols, 1))
        n = len(rows)
        start = 2
        for i in range(2, n + 2):
            cur = rows[i - 2]
            nxt = rows[i - 1] if i - 1 < n else None
            same = (nxt is not None and nxt.get('reg') == cur.get('reg')
                    and nxt.get('block') == cur.get('block')
                    and nxt.get('addr') == cur.get('addr'))
            if not same:
                if i > start:
                    for k in MERGE_KEYS:
                        if k in idx:
                            ws.merge_cells(start_row=start, end_row=i,
                                           start_column=idx[k], end_column=idx[k])
                for k in ('block', 'reg', 'addr'):
                    if k in idx:
                        ws.cell(start, idx[k]).fill = reg_fill
                start = i + 1

    # ---- Summary ----
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'RALF -> Register Map'
    ws['A1'].font = Font(bold=True, size=14)
    info = [('Source RALF', meta['src']), ('Included files', meta['nfiles']),
            ('Top', meta['top']), ('Generated', meta['time']),
            ('Tool', 'ralf2xlsx %s' % VERSION),
            ('Registers', len(ela.reg_rows)), ('Field rows', len(ela.rows)),
            ('Memories', len(ela.mem_rows)), ('Warnings', len(warn_rows)),
            ('Address mode', meta['addr_mode']), ('Arrays', meta['arrays'])]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    r0 = 3 + len(info) + 1
    for j, h in enumerate(['Block', 'Kind', 'Base Addr', 'Bus Bytes', 'Endian', 'Items'], 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, border
    for i, b in enumerate(ela.blk_rows, start=r0 + 1):
        for j, v in enumerate([b['block'], b['kind'], hexs(b['base']),
                               b['bus_bytes'], b['endian'], b['items']], 1):
            c = ws.cell(i, j, v)
            c.border = border
            c.alignment = left if j == 1 else center
    for j, w in enumerate([40, 10, 14, 11, 10, 8], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions['B'].width = 46

    dump(wb.create_sheet('RegisterMap'), ela.rows, REG_COLS, merge=True)
    dump(wb.create_sheet('Registers'), ela.reg_rows, REGSUM_COLS)

    if ela.mem_rows:
        dump(wb.create_sheet('Memories'), ela.mem_rows, MEM_COLS)

    enum_rows, seen = [], set()
    for r in ela.rows:
        key = (r['block'], r['reg'], r['field'])
        if r.get('enum') and key not in seen:
            seen.add(key)
            for item in r['enum'].split(', '):
                if '=' in item:
                    nm, vv = item.split('=', 1)
                    enum_rows.append(dict(block=r['block'], reg=r['reg'],
                                          field=r['field'], name=nm, value=vv))
    if enum_rows:
        dump(wb.create_sheet('Enums'), enum_rows, ENUM_COLS)

    dump(wb.create_sheet('Warnings'), warn_rows, WARN_COLS,
         fill=warn_fill if warn_rows else None)

    if sheet_per_block:
        groups = OrderedDict()
        for r in ela.rows:
            groups.setdefault(r['block'], []).append(r)
        used = set(s.title for s in wb.worksheets)
        for bname, rows in groups.items():
            t = re.sub(r'[\\/*?:\[\]]', '_', bname)[-31:]
            base_t, k = t, 1
            while t in used:
                t = ('%s_%d' % (base_t, k))[-31:]
                k += 1
            used.add(t)
            dump(wb.create_sheet(t), rows, REG_COLS, merge=True)

    wb.save(path)


def write_csv(path, rows, cols):
    import csv
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow([h for _, h, _ in cols])
        for r in rows:
            w.writerow([r.get(k, '') for k, _, _ in cols])


def dump_text(ela, out=sys.stdout):
    """纯文本寄存器表, 方便 diff 两个版本的 RALF"""
    by_reg = OrderedDict()
    for r in ela.rows:
        by_reg.setdefault((r['block'], r['reg'], r['addr']), []).append(r)
    cur_blk = None
    for (blk, reg, addr), frows in by_reg.items():
        if blk != cur_blk:
            out.write('\n=== %s ===\n' % blk)
            cur_blk = blk
        out.write('REG %12s  %-32s %sb\n' % (addr or '-', reg, frows[0]['width']))
        for f in frows:
            if f['field'] == '-':
                continue
            out.write('    %10s %-26s %-6s reset=%s%s\n' %
                      (f['bits'], f['field'], f['access'], f['reset'],
                       ('  ' + f['enum']) if f['enum'] else ''))
    if ela.mem_rows:
        out.write('\n=== memories ===\n')
        for m in ela.mem_rows:
            out.write('MEM %12s  %-32s %s x %sb  %s\n' %
                      (m['addr'] or '-', '%s.%s' % (m['block'], m['name']),
                       m['size'], m['bits'], m['access']))


# ============================================================================
# main
# ============================================================================

def main(argv=None):
    ap = argparse.ArgumentParser(
        description='把 Synopsys RALF 寄存器描述文件转换成 Excel 表格',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='例:\n'
               '  python ralf2xlsx.py top.ralf -t soc -o regmap.xlsx -s\n'
               '  python ralf2xlsx.py top.ralf --check --strict     # CI 里只做检查\n'
               '  python ralf2xlsx.py top.ralf --dump > new.txt     # 文本化后 diff')
    ap.add_argument('ralf', help='顶层 .ralf 文件')
    ap.add_argument('-t', '--top', action='append', default=[],
                    help='顶层 block/system 名, 可多次; 不给则展开所有顶层 system(没有则所有 block)')
    ap.add_argument('-o', '--out', help='输出 .xlsx (缺省: <ralf>.xlsx)')
    ap.add_argument('-I', '--incdir', action='append', default=[], help='source 搜索目录, 可多次')
    ap.add_argument('-s', '--sheet-per-block', action='store_true',
                    help='每个 block 额外生成一张 sheet')
    ap.add_argument('--csv', help='同时输出一份 CSV (字段表)')
    ap.add_argument('--base', default='0', help='顶层基地址 (缺省 0), 支持 0x / Verilog 格式')
    ap.add_argument('--addr-mode', choices=('byte', 'word'), default='byte',
                    help='RALF 里 @offset 的含义: byte=字节偏移(缺省, 对应 ralgen -B / byte addressing); '
                         'word=地址单元个数(= 所在 block 的 bytes), 对应 ralgen 默认地址语义')
    ap.add_argument('--sort', choices=('source', 'addr'), default='source',
                    help='行顺序: source=按 RALF 出现顺序(缺省), addr=按地址')
    ap.add_argument('--no-expand-arrays', action='store_true',
                    help='数组只出一行 (name[lo:hi]), 大数组不撑爆表格')
    ap.add_argument('--check', action='store_true', help='只解析+展开, 不生成 Excel')
    ap.add_argument('--dump', action='store_true', help='把展开后的寄存器表打成纯文本')
    ap.add_argument('--strict', action='store_true', help='有 warning 即失败 (退出码 3)')
    ap.add_argument('-q', '--quiet', action='store_true', help='不打印 warning')
    args = ap.parse_args(argv)

    if not os.path.isfile(args.ralf):
        ap.error('文件不存在: %s' % args.ralf)

    p = Parser(incdirs=args.incdir)
    try:
        p.parse_file(args.ralf)
    except (ParseError, OSError) as e:
        sys.stderr.write('RALF 解析失败: %s\n' % e)
        return 2

    tops = list(args.top)
    if not tops:
        systems = [n for k, n in p.order if k == 'system']
        blocks = [n for k, n in p.order if k == 'block']
        tops = systems or blocks
        if not tops:
            sys.stderr.write('RALF 里没有 block/system, 无法确定顶层\n')
            return 2
        if len(tops) > 1:
            sys.stderr.write('[i] 未指定 -t, 展开全部 %d 个顶层: %s\n'
                             % (len(tops), ', '.join(tops)))

    base, _ = parse_num(args.base)
    ela = Elaborator(p, addr_mode=args.addr_mode,
                     expand_arrays=not args.no_expand_arrays)
    for t in tops:
        topdef = p.defs['system'].get(t) or p.defs['block'].get(t)
        if topdef is None:
            sys.stderr.write('找不到 top: %s (可用: %s)\n' %
                             (t, ', '.join('%s %s' % (k, n) for k, n in p.order)))
            return 2
        ela.walk_block(topdef, topdef.name, base or 0, None)

    if args.sort == 'addr':
        ela.rows.sort(key=lambda r: (r['block'], r['_addr'], r['reg'],
                                     r['lsb'] if r['lsb'] != '' else 0))
        ela.reg_rows.sort(key=lambda r: (r['block'], r['_addr'], r['reg']))

    warn_rows = ([dict(kind='parse', where=w, msg=m) for w, m in p.warnings] +
                 [dict(kind='elab', where='', msg=m) for m in ela.warnings])

    if args.dump:
        dump_text(ela)

    if not args.quiet:
        for w in warn_rows:
            sys.stderr.write('[warn] %s %s\n' % (('%s:' % w['where']) if w['where'] else '',
                                                 w['msg']))

    sys.stderr.write('files %d | registers %d | fields %d | memories %d | warnings %d\n'
                     % (len(p.files), len(ela.reg_rows), len(ela.rows),
                        len(ela.mem_rows), len(warn_rows)))

    if args.strict and warn_rows:
        sys.stderr.write('ERROR: --strict 下有 %d 条 warning, 未生成输出\n' % len(warn_rows))
        return 3

    if args.check:
        return 0

    out = args.out or (os.path.splitext(args.ralf)[0] + '.xlsx')
    meta = dict(src=os.path.abspath(args.ralf), nfiles=len(p.files),
                top=', '.join(tops),
                time=datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                addr_mode='%s offsets (%s), base=%s' %
                          (args.addr_mode,
                           'ralgen -B / byte addressing' if args.addr_mode == 'byte'
                           else 'ralgen default / address units',
                           hexs(base or 0)),
                arrays='collapsed' if args.no_expand_arrays else 'expanded')
    try:
        write_xlsx(out, ela, meta, warn_rows, sheet_per_block=args.sheet_per_block)
    except ImportError:
        sys.stderr.write('ERROR: 需要 openpyxl, 请先 pip install openpyxl\n')
        return 2
    if args.csv:
        write_csv(args.csv, ela.rows, REG_COLS)

    print('[ok] %s' % out)
    return 0


if __name__ == '__main__':
    sys.exit(main())
